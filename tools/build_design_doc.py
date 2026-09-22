import json
import os
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_NODE = Path(r"C:\Users\natha\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\bin\node.exe")


OUT = r"C:\starship\gacha-system\星界之律遊戲設定總覽.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=None, color="000000"):
    run.font.name = "Microsoft JhengHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, size=10.5, color="20242C", bold=False, space_after=5, line=1.25):
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def add_text(doc, text, size=10.5, color="20242C", bold=False, after=5, align=None):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.3
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    size = {1: 18, 2: 13, 3: 11}[level]
    set_run_font(run, size=size, bold=True, color="000000")
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_bullets(doc, items, level=0):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        run = paragraph.add_run(item)
        set_run_font(run, size=10, color="20242C")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.2


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        run = paragraph.add_run(item)
        set_run_font(run, size=10, color="20242C")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.2


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    tr_pr = header._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "193B63")
        set_cell_borders(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(value)
        set_run_font(run, size=9.2, bold=True, color="FFFFFF")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(row):
            cell = cells[col_index]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "F4F7FB" if row_index % 2 else "FFFFFF")
            set_cell_borders(cell)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(str(value))
            set_run_font(run, size=9, color="20242C")
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.15
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_page_break(doc):
    doc.add_page_break()


def load_character_catalog():
    node = Path(os.environ.get("STARSHIP_NODE", str(DEFAULT_NODE)))
    if not node.exists():
        node = Path("node")
    script = "const d=require('./src/data.js'); process.stdout.write(JSON.stringify(Object.values(d.cards)));"
    result = subprocess.run([str(node), "-e", script], cwd=ROOT, check=True, capture_output=True, text=True, encoding="utf-8")
    return sorted(json.loads(result.stdout), key=lambda card: (float(card.get("releaseVersion", 0)), card.get("id", "")))


def add_character_portrait_catalog(doc, portrait_dir):
    cards = load_character_catalog()
    add_page_break(doc)
    add_heading(doc, "角色完整立繪圖鑑", 1)
    add_text(doc, "以下圖鑑與 src/data.js 的 cards 同步，收錄 1.0–5.5 全部角色。每張圖都使用完整立繪，並在圖下固定標示角色名稱、星級、元素、版本與目前開放狀態；後續版本角色雖已建檔，仍會以鎖定狀態保留。", size=10.5, after=8)
    table = doc.add_table(rows=0, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, card in enumerate(cards):
        if index % 3 == 0:
            cells = table.add_row().cells
            for cell in cells:
                cell.width = Inches(2.2)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_shading(cell, "F4F7FB")
                set_cell_borders(cell)
                set_cell_margins(cell, top=120, start=100, bottom=120, end=100)
        cell = table.rows[-1].cells[index % 3]
        cell.text = ""
        image_path = Path(portrait_dir) / (str(card.get("id")) + ".png")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if image_path.exists():
            # Keep the final row together with its captions when Word paginates
            # the gallery.  The smaller width still leaves the portrait large
            # enough to read while avoiding an orphaned caption on a new page.
            paragraph.add_run().add_picture(str(image_path), width=Inches(1.62))
        else:
            fallback = paragraph.add_run("立繪待補")
            set_run_font(fallback, size=9, color="8A3D3D", bold=True)
        paragraph.paragraph_format.space_after = Pt(3)
        caption = cell.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name = caption.add_run(str(card.get("name", "")))
        set_run_font(name, size=10, color="000000", bold=True)
        meta = cell.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        version = str(card.get("releaseVersion", ""))
        status = "玩家開放" if float(version) <= 2.5 else "建檔鎖定"
        planned = card.get("plannedGachaVersion")
        planned_label = str(planned) if planned else "後續再議"
        meta_run = meta.add_run("★" * int(card.get("rarity", 0)) + "  " + str(card.get("element", "")) + "元素  ·  故事 " + version + "  ·  卡池 " + planned_label + "  ·  " + status)
        set_run_font(meta_run, size=8.2, color="3B5C82")
        meta.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_future_character_relationships(doc):
    cards = [card for card in load_character_catalog() if float(card.get("releaseVersion", 0)) >= 3]
    add_heading(doc, "後續角色與瑟蕾雅的關係索引", 2)
    add_text(doc, "故事初登場和預計卡池刻意分開。這份索引記錄每位角色在瑟蕾雅旅程中的情感功能，供主線、合併支線與角色故事共同使用；未進卡池的角色仍然是完整的故事角色。", size=10, after=6)
    rows = []
    for card in cards:
        planned = card.get("plannedGachaVersion") or "後續再議"
        rows.append([
            card.get("name", ""),
            f"故事 {card.get('releaseVersion', '')}／卡池 {planned}",
            card.get("storyRelationship", "") or "待補充"
        ])
    add_table(doc, ["角色", "登場／卡池", "關係與情感作用"], rows, widths=[1.0, 1.25, 4.35])


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft JhengHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0x20, 0x24, 0x2C)
    for style_name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        style = styles[style_name]
        style.font.name = "Microsoft JhengHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.color.rgb = RGBColor(0, 0, 0)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("星界之律 2.0–2.5  設定總覽")
    set_run_font(footer_run, size=8, color="6D7684")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run("星界之律遊戲設定總覽")
    set_run_font(title_run, size=25, bold=True, color="000000")
    title.paragraph_format.space_after = Pt(4)
    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run("玩家版本 2.0–2.5  規則 內容與後續版本規劃")
    set_run_font(subtitle_run, size=11, bold=True, color="3B5C82")
    subtitle.paragraph_format.space_after = Pt(12)
    add_text(doc, "本文件是目前遊戲程式的單一規則總覽，供製作、測試、部署與後續版本維護使用。1.0 的瑟蕾雅與獸靈之村開場保留，1.1 起的主線、合併支線與角色登場順序由 story-replan-v1 統一重編；現行玩家內容開放至 2.5，3.0–5.5 已整理完成但維持鎖定。除抽卡、劇情、角色培養與戰鬥外，星海迷航和星伴培育也已納入同一套帳號保存規則。任何數值或開放狀態調整，都應同步檢查 src/data.js、玩家介面與測試。", size=10.8, after=10)

    add_heading(doc, "版本摘要", 1)
    add_table(doc, ["項目", "目前設定", "狀態"], [
        ["玩家入口", "先登入，再進入星界之律大廳", "已完成"],
        ["劇情", "主線與支線 1.0–2.5，共 18 章、72 幕；1.1 起採故事先行重編", "開放"],
        ["角色卡池", "1.0–1.5 舊限定池、2.0–2.5 潮眼限定池、1.0–2.5 常駐池", "開放"],
        ["星界試煉", "30 關，最多 4 人，自走棋式自動戰鬥", "本次更新重置"],
        ["Boss 突破", "6 種素材 Boss 分成 Lv.1–3 三檔；Lv.3 給更豐富獎勵，通用印記避免卡關", "新增玩法"],
        ["星港委託", "3 份短篇戰鬥任務，每版本各完成一次", "新增玩法"],
        ["星海迷航", "3 條航線、事件選擇、普通／隱藏／特殊結局與季裝", "新增玩法"],
        ["星伴培育", "寵物等級、裝扮、特效、公開展示與玩家評分", "新增玩法"],
        ["後續內容", "3.0–5.5 劇情、角色、敵人與數值", "已建檔並鎖定"],
    ], widths=[1.35, 4.15, 1.1])

    add_heading(doc, "大更新獎勵", 2)
    add_text(doc, "版本標記為 2.0-2.5。伺服器在玩家建立帳號或登入時檢查 updateRewards.claimedVersions，只要該帳號尚未領取本版本，就加入 3,200 星砂並寫入領取時間。新帳號與舊帳號使用同一套邏輯，因此不會因為建立時間不同而漏領或重複領取。試煉獎勵次數與 Boss 挑戰紀錄按版本重置，但玩家角色、等級、命座、突破狀態與已取得的突破材料不重置。", after=4)
    add_bullets(doc, [
        "試煉版本、清除紀錄、最佳關卡、每關獎勵次數和上次戰報在本次更新一併重置。",
        "星港委託的完成紀錄也以 2.0-2.5 為版本鍵，下次版本更新時清空。",
        "Boss 每個版本重新計算成功獎勵次數；已取得的材料與角色突破狀態跨版本保留。",
        "前端本機存檔模式也執行相同的一次性檢查；正式玩家資料以後端為準。",
    ])

    add_page_break(doc)
    add_heading(doc, "抽卡規則", 1)
    add_text(doc, "抽卡核心由 src/gacha.js 統一提供給瀏覽器、Node 伺服器與測試使用。十連只是逐格執行十次單抽判定，不會把保底或 4★ 結果藏到最後一格。", after=7)
    add_table(doc, ["判定位置", "4★機率", "說明"], [
        ["第 1–20 抽", "0%", "前 20 抽不會提早出現 4★"],
        ["第 21 抽", "10%", "開始進入提升區間"],
        ["第 22–49 抽", "每抽 +4 個百分點", "第 49 抽最多顯示 99%"],
        ["第 50 抽", "100%", "硬保底，必定 4★"],
        ["4★內部判定", "精選 55%／其他 45%", "僅限定池使用，歪掉後下一張 4★ 保證精選"],
        ["非 4★結果", "三星 20%／一般回響 80%", "三星直接從該卡池的文件角色清單抽取"],
        ["額外星砂", "8%", "單格另掉落 40 星砂，不改變角色判定"],
    ], widths=[1.35, 1.55, 3.7])

    add_heading(doc, "目前卡池", 2)
    add_table(doc, ["卡池", "精選與範圍", "保底計數"], [
        ["限定 1.0–1.5", "文件 1.0–1.5 的 4★可選一名；三星為 1.0–1.5", "limited"],
        ["限定 2.0–2.5", "璃珊、曜澤、伊芙琳、澪歌、菲芮、諾芮亞、奧薇拉可選一名；三星為 1.0–2.5", "與限定共用 limited"],
        ["常駐回音", "目前已開放的 1.0–2.5 角色；沒有精選", "standard"],
        ["復刻池", "尚未開放", "不顯示於玩家選單"],
    ], widths=[1.45, 4.15, 1.0])

    add_heading(doc, "資源用途", 2)
    add_table(doc, ["資源", "取得方式", "用途"], [
        ["星砂", "大更新、劇情、試煉、委託、抽卡補償與少量抽卡掉落", "單抽 160，十連 1,600"],
        ["星痕", "重複 4★", "10 枚兌換限定池精選 4★"],
        ["回響粉", "一般回響結果", "抽卡系統內的非角色回響資源"],
        ["角色經驗", "試煉、委託、重複 3★等", "角色升級，不與星砂共用"],
        ["突破材料", "Boss 各等級勝利", "指定材料或星界通用突破印記；80 等突破後升到 90 等"],
        ["個人命座晶核", "重複取得該名角色", "只能提升同一角色命座，每次 1 枚"],
        ["寵物資源", "星伴探索、培育、玩家展示評分", "只用於寵物等級、裝扮與公開展示"],
    ], widths=[1.5, 3.0, 2.1])

    add_heading(doc, "養成經濟檢查", 2)
    add_text(doc, "這次調整前先用現行規則跑完整資源盤點，並把劇情的一次性獎勵與試煉／Boss 每版本最多 10 次的可重複獎勵分開計算。角色經驗成本不靠角色面板膨脹來解決，而是讓玩家有穩定、可預期的遊玩來源。", after=5)
    add_table(doc, ["目標", "需求", "檢查結果"], [
        ["單名四星 1→70", "79,902 角色經驗", "可由前期劇情與試煉逐步取得"],
        ["單名四星 1→80", "104,122 角色經驗", "80 等突破前可完成"],
        ["單名四星 1→90", "131,542 角色經驗", "需完成 80 等突破並使用 Boss 材料"],
        ["四名四星 1→90", "526,168 角色經驗", "目前版本列出的全部來源合計 537,300，餘 11,132"],
    ], widths=[1.55, 1.8, 3.25])
    add_bullets(doc, [
        "開放劇情 72 幕是一次性獎勵：每幕 +100 星砂、+650 角色經驗，重看不會重複領取。",
        "星界試煉 30 關每關每版本最多成功領 10 次：每次 +50 星砂、+1,500 角色經驗；保留挑戰次數是為了維持主要可重複的角色經驗來源。",
        "六個 Boss 各 10 次、分布於 Lv.1–3 三檔，合計 +32,400 角色經驗、80 枚專屬材料與 60 枚通用印記；星港委託、星海迷航、新手教學和初始資源合計再提供 8,100 經驗。",
        "完整使用當期玩法上限可取得 28,260 星砂：176 次單抽，並剩 100 星砂；抽卡重複角色帶來的經驗與每格 8% 隨機星砂未列入安全估算。",
    ])
    add_heading(doc, "後續角色天賦預留", 2)
    add_text(doc, "3.0 之後才開放角色天賦。每名角色固定三條分支：「攻擊手段」「定位專精」「界痕共鳴」，各 5 級；使用獨立的專屬天賦手冊，不消耗星砂或突破材料。每條分支上限約 4–6%，三條合計受 10% 低幅度上限約束，技能效果以功能提升為主，不會把角色面板無限乘高。", after=4)

    add_page_break(doc)
    add_heading(doc, "劇情內容", 1)
    add_text(doc, "玩家進入劇情選單後可以切換主線與支線。章節清單只讓 1.0–2.5 可選；3.0–5.5 仍在同一份檔案中顯示為已建檔但鎖定。1.0 保留原始開場，1.1 起採 story-replan-v1；選定章節後，當幕內容與本章所有幕次會同時顯示，避免只看到劇情標題。", after=7)
    add_table(doc, ["版本", "主線", "支線", "開放狀態"], [
        ["1.0", "界痕初響", "獸靈村的回覆", "開放"],
        ["1.1", "聽見的人", "—", "開放／重編"],
        ["1.2", "橋下的遺構", "橋下的交班", "開放／重編"],
        ["1.3", "潮線之外", "—", "開放／重編"],
        ["1.4", "鐘庭的半小時", "門外的半小時", "開放／重編"],
        ["1.5", "空白頁的第一個人", "—", "開放／重編"],
        ["2.0", "沒有地址的書庫", "藍燈不滅", "開放／重編"],
        ["2.1", "鏡潮島的第二句", "—", "開放／重編"],
        ["2.2", "深潮測線", "空船的乘客", "開放／重編"],
        ["2.3", "雲脊的改道路", "—", "開放／重編"],
        ["2.4", "霧鏡議庭", "見證人的空白", "開放／重編"],
        ["2.5", "潮眼之外", "—", "開放／重編"],
    ], widths=[.8, 2.2, 2.2, 1.4])
    add_heading(doc, "劇情獎勵與角色里程碑", 2)
    add_bullets(doc, [
        "每幕第一次完成時加入 100 星砂與 650 角色經驗，完成紀錄以 chapterId:sceneId 儲存，重看不會重複領取。",
        "新帳號自動取得瑟蕾雅。完成 1.0 主線任意一幕後，可在大廳自選雷恩或莉亞。",
        "通關星界試煉第 10 關後，再開啟一次雷恩或莉亞自選，第二份獎勵優先顯示尚未取得的角色。",
    ])

    add_page_break(doc)
    add_heading(doc, "角色培養", 1)
    add_text(doc, "角色頁先顯示已取得角色，再以完整立繪、名稱、星級、元素、戰鬥定位與成長按鈕提供詳細培養。立繪比例使用 contain，名稱、星級與元素以固定標籤疊加，避免瑟蕾雅等角色在手機畫面被裁切。角色戰力成長依實際基礎面板分成標準帶與平衡帶，不以性別決定倍率，讓低基礎的功能型四星在 70–90 等仍有追趕空間。", after=7)
    add_table(doc, ["項目", "三星角色", "四星角色"], [
        ["基礎數值", "整體較低，靠治療、速度或特定功能補位", "生命、攻擊、防禦與整體戰力基準較高"],
        ["升級倍率", "每級 3% 主要屬性、2.2% 防禦、0.9% 速度", "標準每級 4%／3%／1.2%；低基礎戰力 4★進入平衡帶後為 5%／3.6%／1.4%"],
        ["命座倍率", "每命主要屬性 +40%、防禦 +13%、速度 +2.2%、技能倍率 +2%；滿命滿等約接近一般 4★55 等", "每命主要屬性 +12%、防禦 +8%、速度 +2.5%、技能倍率 +1.8%；滿命在高階戰鬥有明顯回饋"],
        ["升級成本", "起始 45 角色經驗，每級增加 18", "起始 70 角色經驗，每級增加 32"],
        ["協同價值", "便宜、快速或提供治療與特殊定位", "輸出、重裝、支援與控制組合較完整"],
    ], widths=[1.35, 2.7, 2.55])
    add_heading(doc, "等級上限與突破", 2)
    add_bullets(doc, [
        "現行角色等級上限為 90；角色可先升到 80 等，達到 80 等後必須取得突破材料才能繼續升級。",
        "每名角色有偏好的 Boss 與專屬材料：三星角色每次突破需要 3 枚，四星角色需要 4 枚；指定材料不足時可用星界通用突破印記補足。",
        "100 等只完成規則與資料預留，本次不在玩家介面開放，也不會被當成目前等級上限。",
    ])
    add_heading(doc, "命座規則", 2)
    add_bullets(doc, [
        "初次取得角色為 0 命；後續每取得一次相同角色，命座直接增加 1，並留下 1 枚該角色專用命座晶核。",
        "例如莉亞持有數量為 5，代表初次取得加上 4 次重複，應為 4 命與 4 枚個人晶核；舊存檔登入時會自動補回。",
        "重複角色已在抽卡結算時自動增加命座；培養頁不再放置重複的手動提升按鈕，個人晶核保留作為存檔相容與後續系統資源。",
        "抽卡不再產生全域共鳴晶核；命座只由重複角色自動增加，該角色專用晶核僅作為個人命座資源與舊存檔相容資料保留。",
    ])

    portrait_dir = os.environ.get("STARSHIP_DOCX_PORTRAIT_DIR")
    if not portrait_dir:
        raise RuntimeError("建立設定文件前必須指定 STARSHIP_DOCX_PORTRAIT_DIR，避免輸出沒有角色立繪的文件")
    add_character_portrait_catalog(doc, portrait_dir)

    add_page_break(doc)
    add_heading(doc, "星界試煉", 1)
    add_text(doc, "星界試煉採自走棋式自動戰鬥。玩家只決定最多四人的隊伍，戰鬥依速度、技能冷卻、敵方特性和隊伍協同自動進行。推薦戰力不是絕對門檻，系統保留低戰力靠配合取勝以及高戰力因編隊失誤而失敗的可能。", after=7)
    add_table(doc, ["規則", "內容"], [
        ["關卡數量", "30 關，逐關解鎖；第 30 關為終局關卡"],
        ["隊伍上限", "一次最多 4 名已取得角色"],
        ["版本次數", "每版本每一關最多成功領獎 10 次"],
        ["成功獎勵", "50 星砂、1,500 角色經驗"],
        ["失敗處理", "保留編隊與戰報，不扣獎勵次數"],
        ["更新處理", "2.0–2.5 更新已重置清除紀錄、最佳關卡、次數與上次戰報"],
        ["敵人資訊", "每關展示敵人圖片、數量、HP、攻擊、防禦、速度與威脅值"],
    ], widths=[1.35, 5.25])
    add_heading(doc, "編隊思路", 2)
    add_bullets(doc, [
        "重裝或守衛承受傷害，治療與修復角色維持隊伍，支援與指揮角色提高整體效率。",
        "測量、仲裁與編譯等角色可降低敵方防禦或清除增益；射手、斥候與爆發角色負責集中擊破。",
        "環境與敵方特性會改變答案，例如治療降低、敵方護盾、噪音延遲技能或高速獵襲。",
        "後期第 21–30 關已重新校準敵方攻擊曲線與首領機制：不再只靠高生命拖長戰鬥，也不把傷害拉到只能靠滿命隊伍硬扛；第 30 關以 12,800 推薦戰力、620／690 基礎攻擊的護衛與王座作為終局基準。",
        "戰鬥演算保護上限為 120 回合；超過上限只會回傳 timeout，不會誤判為通關或直接扣除成功獎勵次數。",
    ])

    add_page_break(doc)
    add_heading(doc, "Boss 與 80 等突破", 1)
    add_text(doc, "Boss 是角色 80 等後的專屬突破玩法。玩家在 Boss 頁選擇首領、查看敵人與環境情報，再從已取得角色中派出最多 4 人。六種素材 Boss 分成 Lv.1–3 三個獎勵檔位，每個檔位有兩個 Boss；Boss 等級不是角色能否突破的硬門檻。每名角色仍會顯示偏好的專屬材料來源，但所有 Boss 都會掉落可替代指定材料的星界通用突破印記。", after=7)
    add_table(doc, ["Boss", "等級", "獎勵檔位", "材料用途"], [
        ["星序守望者", "Lv.1", "1 專屬材料、420 經驗、1 通用印記", "星序碎晶；通用印記可替代"],
        ["潮眼書庫獸", "Lv.1", "1 專屬材料、420 經驗、1 通用印記", "潮眼晶核；通用印記可替代"],
        ["逆時守鐘人", "Lv.2", "1 專屬材料、540 經驗、1 通用印記", "逆時鐘核；通用印記可替代"],
        ["鍛路熔殼王", "Lv.2", "1 專屬材料、540 經驗、1 通用印記", "熱管熔核；通用印記可替代"],
        ["風廊獵王", "Lv.3", "2 專屬材料、660 經驗、1 通用印記", "風標獵核；通用印記可替代"],
        ["霧鏡裁定核", "Lv.3", "2 專屬材料、660 經驗、1 通用印記", "霧鏡映核；通用印記可替代"],
    ], widths=[1.45, .75, 2.5, 2.0])
    add_bullets(doc, [
        "每個 Boss 每版本最多成功領取 10 次獎勵；Lv.3 的獎勵較多，但不是突破的唯一道路，失敗不扣成功獎勵次數。",
        "三星角色突破需要 3 枚材料，四星角色突破需要 4 枚材料；指定材料不足時可以用星界通用突破印記補足。",
        "Boss 的版本挑戰紀錄、編隊與上次戰報會在版本更新時重置；突破材料、角色等級和突破狀態保留在玩家帳號。",
        "Boss 戰同樣採自走棋式自動戰鬥，參考戰力只作提示；角色技能、元素、速度與隊伍協同會影響勝負。",
    ])

    add_page_break(doc)
    add_heading(doc, "星海迷航", 1)
    add_text(doc, "星海迷航是獨立於主線的短局航程。玩家從三條航線中選擇一條，在戰鬥、事件、休整、商店和終端節點之間前進；臨時增益與碎片只在本次航程生效，不會污染角色與主線存檔。", after=7)
    add_table(doc, ["系統", "內容", "獎勵"], [
        ["回聲航線", "收集回聲並調和不同星級隊伍", "普通或特殊結局"],
        ["漂流商路", "用碎片購買臨時增益，安排終局爆發", "普通或隱藏結局"],
        ["無名檔案線", "讀取空白頁、開啟檔案門", "隱藏結局"],
        ["版本獎勵", "普通、隱藏、特殊結局各版本領一次", "星砂、角色經驗、通用資源"],
        ["季節裝扮", "特殊結局完成後解鎖當期一名四星裝扮", "2.0–2.5：梅芙「流光檔案裝」"],
    ], widths=[1.3, 3.75, 1.65])
    add_bullets(doc, [
        "特殊結局需要三星與四星混編並完成協鳴條件；隱藏結局需要檔案線索或足夠回聲。",
        "航程失敗可以重新開始；版本更新只重置航程、結局領取紀錄與臨時進度，不重置角色。",
    ])

    add_heading(doc, "星伴培育", 1)
    add_text(doc, "星伴培育是與主線和角色養成分開的寵物收藏系統。玩家可領養、餵食、玩耍、訓練與探索寵物，並替寵物設定裝扮與出場特效。", after=7)
    add_table(doc, ["項目", "規則"], [
        ["獨立資源", "飼料、玩具、星伴代幣，不消耗星砂、角色經驗或突破材料"],
        ["培育", "寵物等級上限 30；餵食、玩耍、訓練提升經驗、親密度、心情與訓練紀錄"],
        ["展示", "可選寵物、裝扮、特效；展示可設為私人或公開"],
        ["社群評分", "其他玩家可瀏覽公開展示；每期每人對同一作品評分一次，評分者 +1 飼料、擁有者 +1 玩具"],
        ["版本處理", "只重置外出探索次數與評分紀錄，寵物等級、收藏與裝扮設定保留"],
    ], widths=[1.45, 5.25])

    add_page_break(doc)
    add_heading(doc, "星港委託", 1)
    add_text(doc, "星港委託是比試煉更短的任務型玩法。每份委託有自己的環境、敵方特性與戰鬥推薦值，完成一次即領取獎勵；不共用試煉的 10 次獎勵次數，版本更新時一併重置。", after=7)
    add_table(doc, ["委託", "區域與特色", "完成獎勵"], [
        ["潮汐書庫抄錄", "潮汐書庫；敵人數量較多，速度隊伍有優勢", "180 星砂、1,200 角色經驗、4 回響粉"],
        ["白帆岬補燈", "白帆岬；防禦較高，重裝或支援較穩定", "380 星砂、1,500 角色經驗"],
        ["鏡潮回收", "鏡潮島；敵方護盾與增益需要清除和控場", "260 星砂、1,800 角色經驗、1 星痕"],
    ], widths=[1.45, 3.25, 1.9])
    add_heading(doc, "可擴充玩法方向", 2)
    add_bullets(doc, [
        "限時活動：以版本主題提供一次性任務、活動商店與額外劇情。",
        "探索地圖：使用角色定位處理採集、護送、解謎和隱藏事件。",
        "角色故事：為已取得角色提供獨立章節與培養素材。",
        "成就與每日任務：追蹤抽卡、培養、戰鬥協同和劇情完成度。",
    ])

    add_page_break(doc)
    add_heading(doc, "後續版本內容", 1)
    add_text(doc, "1.1–5.5 已完成內容設計並保存於資料層，目前不會被現行卡池或玩家劇情入口使用。1.0 的瑟蕾雅、獸靈之村與四小節旋律保留；1.1 起改採 story-replan-v1，先由主線問題決定需要哪些角色，再讓角色以知識、責任、分歧或情感登場。3.0–3.5 從回聲井走到北門終端，4.0–5.5 才接入北境神話意象。每個大版本仍維持五個小版本，但新四星改為 2–3 名；其餘角色保留為劇情夥伴，3.2、4.2、5.3 另保留偶爾推出三星的空間。", after=7)
    add_table(doc, ["版本", "新角色", "劇情主題", "狀態"], [
        ["3.0–3.5", "霽羅 4★、蘿堤亞 4★、洛恩 4★；岑芽 3★", "回聲井 → 四段水路 → 北門終端；先處理回覆與共同治理，再接上神話篇", "3 名新四星／保留澄音、諾嵐、艾斯特於故事"],
        ["4.0–4.5", "奧蕾雅 4★、凱嵐 4★、伊萊拉 4★；塔莉亞 3★", "新曙港 → 第二條律；把北境神話意象寫成可共同修改的規則", "3 名新四星／保留索萊、涅芙、凱爾於故事"],
        ["5.0–5.5", "維斯妲 4★、伊芮恩 4★、達莉雅 4★；妮拉 3★", "根冠 → 深海回信 → 長冬後的新曙；瑟蕾雅拒絕成為唯一中心並完成交班", "3 名新四星／保留布蘭、薩芙、赫爾凡於故事"],
    ], widths=[.7, 2.25, 2.55, 1.1])

    add_future_character_relationships(doc)

    add_heading(doc, "開放前檢查", 2)
    add_bullets(doc, [
        "確認新版本的主線與合併支線各 3 幕，正文可完整顯示，角色名稱和元素 id 都存在。",
        "確認每個大版本只安排 2–3 名新四星；故事角色仍有完整關係、戰鬥資料和獨立立繪，不因未進卡池而消失。",
        "確認更新版本鍵、試煉重置、委託重置和一次性星砂獎勵同步更新。",
        "確認新角色未誤加入 activeCards 或現行卡池，除非版本正式開放。",
    ])

    add_page_break(doc)
    add_heading(doc, "部署與資料保存", 1)
    add_text(doc, "玩家端與管理端使用同一個 Node 網站。正式部署時，玩家開啟根網址，製作者使用同一網址的 /admin.html；不需要玩家另外啟動本機伺服器。", after=7)
    add_table(doc, ["設定", "用途", "必要性"], [
        ["啟動命令", "node serve.mjs，平台會提供 PORT", "必要"],
        ["NODE_ENV", "正式環境設定為 production", "建議"],
        ["STARSHIP_ADMIN_KEY", "管理頁的私密金鑰，需自行設定長密鑰", "必要"],
        ["DATABASE_URL", "Postgres 連線字串，使用 starship_players 表", "推薦"],
        ["STARSHIP_DATA_DIR", "沒有 Postgres 時的 players.json 目錄", "備援"],
        ["持久化儲存", "避免檔案型玩家資料在重啟後消失", "檔案模式必要"],
    ], widths=[1.65, 3.65, 1.3])
    add_heading(doc, "專案入口", 2)
    add_bullets(doc, [
        "玩家頁：/",
        "測試頁：/test.html，測試抽卡存檔和玩家存檔分開。",
        "管理頁：/admin.html，可查詢玩家與執行測試資源調整。",
        "健康檢查：/api/health，正常時回傳 persistence 為 postgres 或 file。",
        "測試命令：node --test test/gacha.test.js test/battle.test.js。",
    ])
    add_text(doc, "文件與程式目前以 2.0–2.5 為玩家版本基準；1.1 起重編版與 3.0–5.5 內容都已完成建檔，但仍需正式版本公告、資產確認與回歸測試後才可開放。", size=10.5, bold=True, color="193B63", after=0)

    doc.core_properties.title = "星界之律遊戲設定總覽"
    doc.core_properties.subject = "玩家版本 2.0–2.5 規則與 5.5 前內容規劃"
    doc.core_properties.author = "星界之律開發文件"
    doc.save(OUT)


if __name__ == "__main__":
    main()
