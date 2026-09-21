from pathlib import Path
import json
import os
import re
import shutil
import subprocess

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "星界之律遊戲設定總覽.docx"
OUTPUT = ROOT / "星界之律遊戲設定總覽-劇情修訂版.docx"


def load_future_story():
    """Read the 3.0–5.5 story directly from the game's data source.

    Keeping the document generator connected to src/data.js prevents the design
    document from drifting away from the text that the game will display when
    these future chapters are opened.
    """
    bundled_node = Path(r"C:\Users\natha\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\bin\node.exe")
    node = Path(os.environ.get("STARSHIP_NODE", "")) if os.environ.get("STARSHIP_NODE") else bundled_node
    if not node.exists():
        node = Path("node")
    export_script = (
        "import('./src/data.js').then(({default:d}) => "
        "process.stdout.write(JSON.stringify({cards:Object.values(d.cards), story:d.storyChapters, "
        "v3:d.version3StoryChapters, v4:d.version4StoryChapters, v5:d.version5StoryChapters, "
        "releasePlan:d.futureCharacterReleasePlan, futureRevision:d.futureStoryRevision})))"
    )
    result = subprocess.run(
        [str(node), "--input-type=module", "-e", export_script],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    return json.loads(result.stdout)


def replace_in_paragraph(paragraph, replacements):
    text = paragraph.text
    new_text = text
    for old, new in replacements:
        new_text = new_text.replace(old, new)
    if new_text != text:
        paragraph.text = new_text


def iter_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)
    return paragraph


def add_story_table(document):
    rows = [
        ("3.0", "回聲之門", "村口的三面旗", "瑟蕾雅回到獸靈之村，第一次把母親的拒絕放在自己的思念之前，留下由村民共同保管的空位。"),
        ("3.1", "誰替她寫回覆", "黑木匣裡的未完成句", "回覆台模仿瑟蕾雅的聲音；她刪除預設聲紋，讓『我不知道』也能被完整保留。"),
        ("3.2", "白榆河的黑星", "白榆河的輪班表", "她熄掉中央主訊號，讓澄音與岑芽接手四段水路，並在河床發現指向鍛路鎮的黑星。"),
        ("3.3", "火把不是王冠", "空白握柄與兩小時後", "洛恩與她把火源、停爐權和撤退路線交給鎮民，救急不再等於坐上唯一的座位。"),
        ("3.4", "北門的母親", "風路守望表", "艾妲的信留下『不要把我帶走』；瑟蕾雅不把拒絕變成追趕任務，改用可回頭測線前進。"),
        ("3.5", "門外先寫信", "把第一頁留給活著的人", "她知道門是為自己開的，仍先寫信、先詢問，留下可拒絕、可撤回、可重談的第一條星界之律。"),
        ("4.0", "新曙港的第一束光", "天文台輪班表", "根系圖以她的第一筆為中心；她讓第一束光照亮退回點，而不是最短航路。"),
        ("4.1", "碎星工坊的熱源", "空爐旁的工具架", "霜核保存她與母親的身世片段；她讓霜火輪值，不讓能力再次集中在自己身上。"),
        ("4.2", "遠望塔的長距離回覆", "信標見習筆記", "虹橋以母親的聲音呼喚她；兩端都保留關閉權，熟悉的聲音不再直接等於邀請。"),
        ("4.3", "白夜航路的記憶", "失效訊息清單", "三條命線呈現三種可能的瑟蕾雅；她保留本人取回記憶的權利，不強迫復原。"),
        ("4.4", "海溝守門人", "四把鑰匙的交班", "根門只認她一人的聲音；她把開門與叫停拆成四把鑰匙，聽見母親要求她不要獨自進門。"),
        ("4.5", "第二條律", "可撤回協議手冊", "長冬是中央系統等待她簽名的結果；她公開拒絕、撤回與重談，北境神話篇正式轉向共同承擔。"),
        ("5.0", "根冠上的第十盞燈", "北境的根名冊", "第十盞燈刻著『不要點亮』；她把九份根系交回各界，不以世界中心的身分追母親。"),
        ("5.1", "霜火雙核", "長夜裡的工具架", "冷核還原她曾主動選擇離開的記憶；她不燒掉痛苦，也不把它交給公共權力。"),
        ("5.2", "虹橋以外的回覆", "橋上不設王座", "虹橋叫出她的原名；她把名字拿回自己手中，不讓名字成為門的鑰匙。"),
        ("5.3", "命線織庭的空白梭", "織線學徒的三次練習", "三條未來都像真的；她不替任何線主剪掉可能，把空白梭交回選擇者。"),
        ("5.4", "深海的回聲守門人", "第四把鑰匙的交班", "艾妲要求她不要獨自進門；瑟蕾雅選擇安全回到水面，帶回一封不必立刻回答的信。"),
        ("5.5", "長冬後的九界新曙", "把神話寫回人手", "終端能讓艾妲回來但會重建唯一中心；她選擇交班，成為可被替換的見證人，收到母親回信。"),
    ]
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_repeat = OxmlElement("w:tblHeader")
    header_repeat.set(qn("w:val"), "true")
    header_properties.append(header_repeat)
    headers = ["版本", "主線", "支線", "本版本的中心轉折"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for version, main, side, twist in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, [version, main, side, twist]):
            cell.text = text
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    return table


def add_future_story_body(document, story):
    """Append the complete 3.0–5.5 main and side-story scenes."""
    document.add_heading("五、3.0–5.5 後續版本完整正文", level=2)
    intro = document.add_paragraph()
    intro.add_run("閱讀定位：").bold = True
    intro.add_run(
        "1.0–2.5 的瑟蕾雅、獸靈之村與既有世界設定是正史根基；3.0–3.5 使用『回聲之門』重製版，"
        "讓新角色先以故事夥伴身分慢慢建立關係，再在合適的大版本進入卡池。新版先處理瑟蕾雅回到村子、面對母親拒絕與走向北門的個人弧線；4.0–5.5 再以北歐神話意象"
        "重新理解根系、命線、霜火、虹橋、深海守門與長冬，但不直接套用既有神名或神話劇本。"
        "每一版都讓瑟蕾雅面對一個看似只能由她決定的中心，再把拒絕、撤回與交班的權利交還給受影響的人。"
        "每個大版本維持五個小版本，但只安排 2–3 名新四星；其他已完成角色保留於主線與支線，"
        "讓瑟蕾雅與每個人的信任、衝突和情感都能有回收。以下正文與 src/data.js 的未開放版本資料同步。"
    )

    for version_group, label in (
        (story.get("v3", []), "3.0–3.5"),
        (story.get("v4", []), "4.0–4.5"),
        (story.get("v5", []), "5.0–5.5"),
    ):
        document.add_heading(f"{label} 版本正文", level=3)
        for chapter in version_group:
            version = chapter.get("version", "")
            title = chapter.get("title", "未命名章節")
            document.add_heading(f"{version} {title}", level=4)
            summary = document.add_paragraph()
            summary.add_run("章節摘要：").bold = True
            summary.add_run(chapter.get("summary", ""))
            for index, scene in enumerate(chapter.get("scenes", []), start=1):
                document.add_heading(f"{index}. {scene.get('title', '未命名幕次')}", level=5)
                document.add_paragraph(scene.get("body", ""))


def add_character_scope(document, story):
    """Show the complete 1.0–5.5 character plan and its live/locked boundary."""
    document.add_heading("角色規劃與開放狀態", level=2)
    document.add_paragraph(
        "角色資料與劇情資料同樣完整建檔至 5.5。玩家目前只會在 1.0–2.5 取得與培養角色；"
        "3.0–5.5 的角色已保存名稱、星級、元素、定位、故事關係與獨立立繪來源。故事初登場版本"
        "不等於卡池版本：每個大版本只安排 2–3 名新四星，其餘角色先以劇情夥伴身分保留，直到"
        "有足夠篇幅完成她們與瑟蕾雅的關係，再於後續版本正式開放。"
    )
    cards = sorted(story.get("cards", []), key=lambda card: (float(card.get("releaseVersion", 0)), card.get("id", "")))
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_repeat = OxmlElement("w:tblHeader")
    header_repeat.set(qn("w:val"), "true")
    header_properties.append(header_repeat)
    headers = ["故事初登場", "預計卡池", "角色", "星級", "元素", "狀態"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for card in cards:
        cells = table.add_row().cells
        version = str(card.get("releaseVersion", ""))
        status = "玩家開放" if float(version) <= 2.5 else ("故事保留／待後續安排" if card.get("plannedGachaVersion") is None else "已建檔／鎖定")
        planned = card.get("plannedGachaVersion") or "後續再議"
        values = [version, planned, card.get("name", ""), "★" * int(card.get("rarity", 0)), card.get("element", ""), status]
        for cell, text in zip(cells, values):
            cell.text = str(text)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    return table


def add_character_release_plan(document, story):
    """Document the slower release cadence separately from story debut order."""
    document.add_heading("大版本角色發佈節奏", level=2)
    document.add_paragraph(
        "每個大版本維持五個小版本，但不再要求每個小版本都推出新四星。"
        "以下是目前的 3.0–5.5 規劃：新四星控制在每個大版本 2–3 名，"
        "已設計但尚未進卡池的角色仍會在主線、合併支線或角色故事中登場，"
        "等關係與個人弧線成熟後再安排，不把他們當成一次性報到角色。"
    )
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["大版本", "新四星（預計卡池）", "三星", "故事保留角色", "敘事重點"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for plan in story.get("releasePlan", []) or []:
        cards = {card.get("id"): card for card in story.get("cards", [])}
        def names(ids):
            return "、".join(cards.get(card_id, {}).get("name", card_id) for card_id in ids) or "無"
        cells = table.add_row().cells
        values = [
            plan.get("majorVersion", ""),
            names(plan.get("fourStarIds", [])),
            names(plan.get("threeStarIds", [])),
            names(plan.get("storyOnlyIds", [])),
            plan.get("focus", "")
        ]
        for cell, text in zip(cells, values):
            cell.text = str(text)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    return table


def add_character_relationships(document, story):
    """Keep the emotional purpose of every future character explicit."""
    document.add_heading("瑟蕾雅與後續角色關係索引", level=2)
    document.add_paragraph(
        "角色可以先在主線或合併支線中出現，再等關係成熟後進入卡池。"
        "本表不是角色簡介，而是每個角色在瑟蕾雅旅程中的情感功能："
        "她們如何互相質疑、照顧、拉回現實，或讓瑟蕾雅學會把選擇交還給別人。"
    )
    cards = sorted(
        [card for card in story.get("cards", []) if float(card.get("releaseVersion", 0)) >= 3],
        key=lambda card: (float(card.get("releaseVersion", 0)), card.get("id", ""))
    )
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["角色", "故事初登場／預計卡池", "與瑟蕾雅的關係與情感作用"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for card in cards:
        cells = table.add_row().cells
        debut = str(card.get("releaseVersion", ""))
        planned = card.get("plannedGachaVersion") or "後續再議"
        values = [
            card.get("name", ""),
            f"故事 {debut}／卡池 {planned}／{'★' * int(card.get('rarity', 0))}",
            card.get("storyRelationship", "") or "待補充"
        ]
        for cell, text in zip(cells, values):
            cell.text = str(text)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    return table


def add_future_story_revision(document, story):
    """Document the active 3.0–3.5 rewrite without deleting the 1.0–2.5 canon."""
    revision = story.get("futureRevision") or {}
    if not revision:
        return
    document.add_heading("3.0–3.5 劇情重製版：回聲之門", level=2)
    document.add_paragraph(
        "這是目前鎖定版本採用的新版橋接稿。1.0–2.5 的瑟蕾雅、獸靈之村、雷恩、莉亞、伊薩爾、"
        "艾妲與潮眼事件維持原始正史；3.0–3.5 改以回到最初的家、面對母親的拒絕、再走向北門為主軸。"
        "原先的未開放正文不刪除，僅作為可回溯素材；遊戲與本文件以本重製版的摘要、導讀與幕次為準。"
    )
    premise = document.add_paragraph()
    premise.add_run("重製核心：").bold = True
    premise.add_run(revision.get("premise", ""))
    for rule in revision.get("designRules", []) or []:
        add_bullet(document, rule)

    chapters = revision.get("chapters", {}) or {}
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["章節", "新版標題", "新版摘要", "幕次"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for chapter_id, chapter in chapters.items():
        cells = table.add_row().cells
        values = [
            chapter_id,
            chapter.get("title", ""),
            chapter.get("summary", ""),
            "、".join(scene.get("title", "") for scene in chapter.get("scenes", [])),
        ]
        for cell, text in zip(cells, values):
            cell.text = str(text)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)


def add_story_index(document, story):
    """Add an authoritative 1.0–5.5 main/side index from the game data."""
    document.add_heading("劇情版本索引（以遊戲資料為準）", level=2)
    document.add_paragraph(
        "本表直接取自 src/data.js 的 storyChapters，保留 1.0–5.5 的主線與支線名稱。"
        "玩家入口目前只開放 1.0–2.5；3.0–5.5 已完成建檔但維持鎖定。"
    )
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_repeat = OxmlElement("w:tblHeader")
    header_repeat.set(qn("w:val"), "true")
    header_properties.append(header_repeat)
    headers = ["版本", "類型", "章節", "狀態"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for chapter in story:
        cells = table.add_row().cells
        version = str(chapter.get("version", ""))
        chapter_type = "主線" if chapter.get("type") == "main" else "支線"
        status = "玩家開放" if float(version) <= 2.5 and chapter.get("releaseOpen") is not False else "已建檔／鎖定"
        for cell, text in zip(cells, [version, chapter_type, chapter.get("title", ""), status]):
            cell.text = str(text)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    return table


def append_story_revision(document):
    document.add_page_break()
    document.add_heading("劇情流暢度校訂與版本主線", level=1)
    p = document.add_paragraph()
    p.add_run("本次校訂目的：").bold = True
    p.add_run("讓主線與支線都能獨立讀懂，又能回到瑟蕾雅的世界核心。主線推進界痕、母親艾妲與星界之律的核心謎團；支線補足角色選擇、地方後果與主線沒有停留的情緒，最後各自留下可追蹤的線索。")

    document.add_heading("一、全劇唯一主軸", level=2)
    add_bullet(document, "瑟蕾雅不是單純的旁觀者或任務接收者；她是界痕系統的預設收件人，也是所有版本爭奪、保護與質疑的中心。")
    add_bullet(document, "主線依序回答：她為何能讀取界痕、艾妲為何被留在彼岸、黑晶與回覆台如何利用她，以及她是否願意成為世界唯一的中心。")
    add_bullet(document, "每個版本都保留三個閱讀節點：開場提出一個清楚問題，中段讓瑟蕾雅的直覺或能力造成代價，終幕給出部分答案並留下下一版更大的問題。")
    add_bullet(document, "3.0–3.5 是『不讓世界只有一個中心』的轉折；4.0–5.5 是以北歐神話意象重新理解根系、命線、霜火、虹橋與終末重建，但不直接套用既有神名或神話劇本。")

    document.add_heading("二、1.0–2.5 正文整理", level=2)
    add_bullet(document, "目前玩家可讀主線與支線共 18 章（12 章主線＋6 篇合併支線）；原 Google 文件長篇正文仍是主要來源，部署時由 src/story-source.js 固定帶入，不要求玩家連線到文件。")
    add_bullet(document, "長篇正文原先有多個『第四幕』實際包含後續時間線，已在遊戲資料層按文件內的幕次、時間線與地點重新分段；完整正文仍保留，舊存檔的既有幕次 id 也保留。")
    add_bullet(document, "2.1 分頁已補回文件中的完整正文，並校對為獨立的四幕內容；新版本先讓玩家看見折光、母親聲音與『熟悉不等於證據』的轉折。")
    add_bullet(document, "早期支線中只有三句佔位文字的章節已補成短篇：保留各角色視角，加入起因、互動、選擇與回扣主線的收束，不再讓支線看起來像空白按鈕。")
    add_bullet(document, "長篇切幕時以原段落的時間線與標題對齊，不把下一幕的標題錯掛到上一幕；完整正文和既有完成紀錄仍保持相容。")
    add_bullet(document, "主線 1.0–2.5 的章節導讀會顯示『瑟蕾雅主線、懸念、收束』，避免玩家只看到章名而不知道本章在整個故事中的作用。")

    document.add_heading("三、主線與支線的閱讀規則", level=2)
    add_bullet(document, "主線一定要改變瑟蕾雅對世界或自己的理解；只有事件發生、沒有她的選擇與代價，不算主線推進。")
    add_bullet(document, "支線不重複主線摘要，而是回答『這件事落到普通人身上後怎麼辦』；支線結尾要補回一個主線線索、情感回音或可實行的規則。")
    add_bullet(document, "對話可保留雷恩、莉亞、QWER 等角色的幽默和摩擦，但每段輕鬆互動都要服務於信任、分工、休息、拒絕或回家的主題。")
    add_bullet(document, "不得把瑟蕾雅的母親、身世或私人地址直接變成所有人的公共獎勵；重要不等於可以被所有人取用。")

    document.add_heading("四、3.0–5.5 版本校訂表", level=2)
    add_story_table(document)

    story = load_future_story()
    add_story_index(document, story.get("story", []))
    add_character_scope(document, story)
    add_character_release_plan(document, story)
    add_future_story_revision(document, story)
    add_character_relationships(document, story)
    add_future_story_body(document, story)

    document.add_heading("六、完成前檢查清單", level=2)
    add_bullet(document, "每一章有可讀的標題、摘要、至少三幕正文、開場問題、中段代價與結尾轉折。")
    add_bullet(document, "每個主線章節的角色標籤與導讀都能找到瑟蕾雅；支線至少說明它如何受她的選擇影響，且不取代支線角色自己的決定。")
    add_bullet(document, "劇情頁同時顯示當幕閱讀區、本章完整正文、幕次導覽與首次完成獎勵；長文可向下捲動，不再只顯示劇情標題。")
    add_bullet(document, "每次版本更新只更新劇情開放與版本玩法，不清除角色、等級、命座、突破、資源或已完成劇情；文件、src/data.js、測試必須一起核對。")
    add_bullet(document, "本次審核結果：現行 1.0–2.5 章節可讀、幕名與正文已對齊，支線已按 2–3 個小版本合併為六篇完整篇章；3.0–5.5 內容已改為完整的主線／合併支線版本稿，並把角色發佈改為每大版本 2–3 名新四星，維持鎖定直到未來資產與版本公告確認後開放。")


def main():
    document = Document(SOURCE)
    replacements = [
        ("主線與支線 1.0–2.5，共 18 章，每章 3 幕", "主線與支線 1.0–2.5，共 18 章（12 章主線＋6 篇合併支線）；正文依幕次導覽"),
        ("主線與支線 1.0–2.5，一共 18 章，每章 3 幕", "主線與支線 1.0–2.5，共 18 章（12 章主線＋6 篇合併支線）；正文依幕次導覽"),
        ("3.0–4.5", "3.0–5.5"),
    ]
    for paragraph in iter_paragraphs(document):
        replace_in_paragraph(paragraph, replacements)
    append_story_revision(document)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
